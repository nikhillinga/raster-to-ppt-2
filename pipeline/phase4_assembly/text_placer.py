"""Phase 4: Text placement."""

from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor

import config
from pipeline.models import DetectedElement, ElementTree

def render_text_lines(element: DetectedElement, target, tree: ElementTree):
    """Place each OCRLine as its own no-wrap text box at its true position."""
    # We must import px_to_emu inside the function to avoid circular imports 
    # if it's imported at the module level.
    from pipeline.phase4_assembly.assembler import px_to_emu
    
    for line in element.ocr_lines:
        if line.is_art or not line.text.strip():
            continue

        txBox = target.shapes.add_textbox(
            Emu(px_to_emu(line.bbox.x, "x", tree.source_image_w, tree.source_image_h)),
            Emu(px_to_emu(line.bbox.y, "y", tree.source_image_w, tree.source_image_h)),
            Emu(px_to_emu(min(int(line.bbox.w * 1.2), tree.source_image_w - line.bbox.x), "x", tree.source_image_w, tree.source_image_h)),
            Emu(px_to_emu(line.bbox.h, "y", tree.source_image_w, tree.source_image_h)),
        )
        tf = txBox.text_frame
        tf.word_wrap = False
        
        # add_textbox creates a paragraph by default
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = line.text
        run.font.name = config.FONT_FAMILY
        run.font.size = Pt(
            max(config.FONT_SIZE_MIN_PT, 
                min(line.font_size_pt, config.FONT_SIZE_MAX_PT))
        )
        run.font.bold = line.bold
        run.font.color.rgb = RGBColor(*line.color_rgb)
